# -*- coding: utf-8 -*-
"""Renderiza las slides de un deck a PNG, para poder MIRARLAS.

POR QUE HACE FALTA RENDERIZAR (2026-09-09)
------------------------------------------
El diff mes-contra-mes (diff_reportes_mensual.py) caza numeros VIEJOS. No caza
un numero que cambio al valor equivocado, ni nada visual. Y los dos errores mas
caros del cierre de Agosto fueron visuales:

  - Las etiquetas del grafico de benchmark mostraban valores de Julio contra una
    tabla de Agosto EN LA MISMA PAGINA.
  - Desalineados verticales heredados del template (factsheet slide 2 shapes
    61/62, pitch book slide 10 shape 36), identicos desde Abril.

Ninguno de los dos deja rastro en el XML: el texto es correcto, lo que esta mal
es donde cae o contra que se lo compara.

EL DETALLE QUE ROMPE EL ATAJO
-----------------------------
No sirve buscar strings en el PDF. Las autoshapes del pitch book se exportan
VECTORIZADAS: en la slide 10 el PDF no devuelve ni "ASSET ALLOCATION" aunque
esta escrito ahi. Buscar texto da falsos negativos. Hay que rasterizar y mirar.
Los charts si conservan texto extraible, pero no se puede depender de eso.

DE DONDE SALE LA IMAGEN
-----------------------
Del PDF que Lucas ya exporta al lado de cada .pptx. Se verifica que exista, que
tenga la misma cantidad de paginas que slides, y que NO sea mas viejo que el
pptx -- un PDF viejo mostraria el deck del mes pasado y la verificacion diria
que todo esta bien. Si no cumple, se aborta con instrucciones en vez de
renderizar algo enganoso.

USO
---
    python scripts/render_slides.py <deck.pptx> [--out DIR] [--dpi 150]
    python scripts/render_slides.py <deck.pptx> --slides 2,10
"""
import argparse
import os
import sys
from datetime import datetime

import fitz  # PyMuPDF
from pptx import Presentation

# 150 dpi: en 16:9 da ~1500x844, suficiente para leer un eje de grafico sin
# que cada slide pese varios MB.
DPI_DEFAULT = 150


def pdf_hermano(pptx_path):
    """El PDF que Lucas exporta al lado del pptx, con el mismo nombre."""
    return os.path.splitext(pptx_path)[0] + ".pdf"


def verificar(pptx_path, pdf_path):
    """Devuelve (n_slides, problemas[]). Vacio = se puede confiar en el PDF."""
    problemas = []
    pres = Presentation(pptx_path)
    n_slides = len(pres.slides)

    if not os.path.exists(pdf_path):
        problemas.append(
            "No existe el PDF %s. Exportarlo desde PowerPoint "
            "(Archivo > Exportar > Crear PDF) y volver a correr."
            % os.path.basename(pdf_path))
        return n_slides, problemas

    doc = fitz.open(pdf_path)
    if doc.page_count != n_slides:
        problemas.append(
            "El PDF tiene %d paginas y el pptx %d slides. Re-exportar el PDF."
            % (doc.page_count, n_slides))

    t_pptx = os.path.getmtime(pptx_path)
    t_pdf = os.path.getmtime(pdf_path)
    if t_pdf < t_pptx:
        problemas.append(
            "El PDF es MAS VIEJO que el pptx (pdf %s, pptx %s). Estaria "
            "verificando el deck de antes de los ultimos cambios. Re-exportarlo."
            % (datetime.fromtimestamp(t_pdf).strftime("%Y-%m-%d %H:%M"),
               datetime.fromtimestamp(t_pptx).strftime("%Y-%m-%d %H:%M")))
    return n_slides, problemas


def render(pptx_path, out_dir, dpi=DPI_DEFAULT, solo=None):
    """Escribe out_dir/slideNN.png. Devuelve la lista de paths."""
    pdf_path = pdf_hermano(pptx_path)
    n_slides, problemas = verificar(pptx_path, pdf_path)
    if problemas:
        for p in problemas:
            print("  ERROR: %s" % p)
        sys.exit(1)

    os.makedirs(out_dir, exist_ok=True)
    doc = fitz.open(pdf_path)
    zoom = dpi / 72.0
    mat = fitz.Matrix(zoom, zoom)

    salidas = []
    for i in range(doc.page_count):
        nro = i + 1
        if solo and nro not in solo:
            continue
        pix = doc.load_page(i).get_pixmap(matrix=mat)
        destino = os.path.join(out_dir, "slide%02d.png" % nro)
        pix.save(destino)
        salidas.append(destino)
        print("  slide %2d -> %s (%dx%d)" % (nro, os.path.basename(destino),
                                             pix.width, pix.height))
    return salidas


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("pptx", help="deck a renderizar")
    ap.add_argument("--out", default=None,
                    help="directorio de salida (default: junto al pptx, en _render/)")
    ap.add_argument("--dpi", type=int, default=DPI_DEFAULT)
    ap.add_argument("--slides", default=None,
                    help="solo estas slides, separadas por coma (ej: 2,10)")
    a = ap.parse_args()

    if not os.path.exists(a.pptx):
        sys.exit("No existe %s" % a.pptx)

    solo = None
    if a.slides:
        solo = {int(x) for x in a.slides.split(",") if x.strip()}

    out = a.out or os.path.join(os.path.dirname(os.path.abspath(a.pptx)), "_render")
    print("Renderizando %s" % os.path.basename(a.pptx))
    print("  desde : %s" % os.path.basename(pdf_hermano(a.pptx)))
    print("  hacia : %s" % out)
    paths = render(a.pptx, out, a.dpi, solo)
    print("\n  %d slide(s) renderizada(s)." % len(paths))


if __name__ == "__main__":
    main()
