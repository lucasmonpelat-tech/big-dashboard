# -*- coding: utf-8 -*-
"""Los dos reportes tienen que decir lo mismo.

POR QUE (2026-09-09)
--------------------
El factsheet y el pitch book se arman por separado y repiten muchos datos del
fondo: ISIN, custodio, minimo de inversion, liquidacion, fees, y las estadisticas
de renta fija. Nada obliga a que coincidan.

En el cierre de Agosto ya paso una version de esto: el factsheet quedo con las
etiquetas del grafico de benchmark en valores de Julio mientras la tabla de la
misma pagina estaba en Agosto. Se detecto en la auditoria final, de casualidad.

Los numeros que YA tienen fuente declarada no necesitan este chequeo: si los dos
decks salen del mismo campo del mismo JSON, coinciden por construccion (eso lo
garantiza check_reportes_vs_datos.py). Este script cubre el resto: los datos del
fondo que se tipean a mano en los dos lados.

COMO MATCHEA
------------
No por shape_id -- son decks distintos, los ids no significan lo mismo. Matchea
por ETIQUETA: busca pares "Etiqueta: valor" en el texto de ambos y compara los
valores de las etiquetas que aparecen en los dos.

Las etiquetas estan en idiomas distintos (el factsheet en castellano, la ficha
tecnica del pitch book en ingles), asi que hay una tabla de equivalencias.

USO
---
    python scripts/check_reportes_cruzados.py <factsheet.pptx> <pitchbook.pptx>

Exit code 1 si algun dato compartido no coincide.
"""
import argparse
import os
import re
import sys
import unicodedata

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402

from diff_reportes_mensual import norm  # noqa: E402

# Etiquetas que significan lo mismo en los dos decks. Clave = nombre canonico.
EQUIVALENCIAS = {
    "isin": ["isin"],
    "moneda": ["moneda", "currency"],
    "custodio": ["custodio", "custodian"],
    "minimo de inversion": ["minimo de inversion", "minimum investment amount",
                            "minimum investment"],
    "liquidacion": ["liquidacion", "settlement"],
    "estructurador": ["estructurador", "programme coordinator"],
    "manager": ["manager", "portfolio manager"],
    "emisor": ["activo", "issuer"],
    "suscripciones": ["suscripciones", "subscriptions"],
    "rescates": ["rescates", "redemptions"],
    "max rescate": ["maximo rescate trimestral", "max redemption"],
    "nombre del fondo": ["fund name"],
}

# Valores que se escriben distinto y significan lo mismo. Se comparan
# normalizados; esto evita falsos positivos por idioma o formato.
SINONIMOS = {
    "t + 2": "t+2",
    "10.000 usd": "10000",
    "$10.000": "10000",
    "$10,000": "10000",
    "10,000 usd": "10000",
    "mensual": "monthly",
    "pro capital": "procapital",
    "lynk capital markets": "lynk capital markets",
}

SEPARADOR = re.compile(r"^(?P<lab>[^:]{2,44}):\s*(?P<val>.+)$")


def canon(s):
    """Minusculas, sin acentos, espacios colapsados."""
    s = unicodedata.normalize("NFKD", (s or "").lower())
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip()


def etiqueta_canonica(lab, val=""):
    """El nombre canonico de una etiqueta. El valor desempata cuando hace falta.

    'Rescates' significa dos cosas distintas segun donde aparezca: en el pitch
    book es la FRECUENCIA ('Redemptions: Monthly') y en el factsheet, en la
    segunda caja, es el TOPE ('Rescates -> Maximo rescate trimestral 5% del
    NAV'). Compararlas entre si da un falso positivo garantizado todos los
    meses, que es peor que no chequear: un reporte que grita en falso se ignora.
    Se desempata por el valor, que es lo unico que distingue los dos casos.
    """
    c = canon(lab)
    if c == "rescates" and "maximo" in canon(val):
        return "max rescate"
    if c == "suscripciones / rescates":
        return "suscripciones"
    for nombre, variantes in EQUIVALENCIAS.items():
        if c in variantes:
            return nombre
    return None


# El mismo hecho escrito en dos idiomas. Se comparan los conceptos, no las
# palabras: "Maximo rescate trimestral 5% del NAV" y "5% of NAV QoQ" dicen
# exactamente lo mismo. Sin esto el chequeo marca una diferencia todos los meses,
# y un reporte que grita en falso se termina ignorando.
CONCEPTOS = {
    "trimestral": "trim", "qoq": "trim", "quarterly": "trim",
    "mensual": "mes", "monthly": "mes",
    "anual": "anio", "annual": "anio", "yearly": "anio",
    "nav": "nav", "usd": "", "del": "", "of": "", "the": "", "maximo": "",
    "max": "", "rescate": "", "redemption": "",
}


def valor_comparable(v):
    c = canon(v)
    c = SINONIMOS.get(c, c)
    # "$10.000" -> "10000" para poder comparar con "10.000 USD"
    solo_num = re.sub(r"[^\d]", "", c)
    if solo_num and len(solo_num) >= 4 and re.fullmatch(r"[\$\d.,\s]*(usd)?", c):
        return solo_num
    # Frase: se queda con los numeros y los conceptos, ordenados. Asi
    # "maximo rescate trimestral 5% del nav" == "5% of nav qoq".
    piezas = set()
    for tok in re.findall(r"[a-z]+|\d+(?:[.,]\d+)?", c):
        if tok.replace(",", ".").replace(".", "").isdigit():
            piezas.add(tok.replace(",", "."))
        else:
            mapeado = CONCEPTOS.get(tok, tok)
            if mapeado:
                piezas.add(mapeado)
    return " ".join(sorted(piezas)) if piezas else c


def _parrafos(pptx):
    """(etiqueta_shape, texto_del_parrafo) preservando las TABULACIONES.

    No se puede usar recolectar(): normaliza los espacios en blanco y ahi se
    pierden los tabs. El factsheet arma su ficha con 'Etiqueta\\tvalor\\tEtiqueta
    \\tvalor' en un mismo parrafo -- colapsado queda 'Activo Lseries DAC
    Estructurador Lynk Capital Markets' y no hay forma de saber donde termina
    una cosa y empieza la otra.
    """
    pres = Presentation(pptx)
    out = []

    def walk(i, shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(i, sh.shapes)
                continue
            if sh.has_table:
                for ri, row in enumerate(sh.table.rows):
                    celdas = [c.text.strip() for c in row.cells]
                    if len(celdas) >= 2 and celdas[0]:
                        out.append(("s%d tabla id=%d f%d" % (i, sh.shape_id, ri),
                                    celdas[0] + "\t" + celdas[1]))
            elif sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs)
                    if t.strip():
                        out.append(("s%d txt id=%d" % (i, sh.shape_id), t))

    for i, s in enumerate(pres.slides, start=1):
        walk(i, s.shapes)
    return out


def hechos(pptx):
    """{etiqueta_canonica: (valor_original, de_donde)} del deck.

    Dos formatos conviven:
      pitch book -> 'Etiqueta: valor', una por linea
      factsheet  -> 'Etiqueta\\tvalor' y a veces dos pares en el mismo parrafo
    """
    out = {}

    def guardar(lab, val, etq):
        nombre = etiqueta_canonica(lab, val)
        if nombre and nombre not in out and val.strip():
            out[nombre] = (norm(val), etq)

    for etq, texto in _parrafos(pptx):
        if "\t" in texto:
            partes = [p for p in texto.split("\t") if p.strip()]
            for i in range(0, len(partes) - 1, 2):
                guardar(partes[i], partes[i + 1], etq)
            continue
        for linea in re.split(r"(?<=[a-z0-9%\)])\s+(?=[A-Z][a-zA-Z ]{2,40}:)", texto):
            m = SEPARADOR.match(linea.strip())
            if m:
                guardar(m.group("lab"), m.group("val"), etq)
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("factsheet")
    ap.add_argument("pitchbook")
    a = ap.parse_args()

    for p in (a.factsheet, a.pitchbook):
        if not os.path.exists(p):
            sys.exit("No existe %s" % p)

    fs, pb = hechos(a.factsheet), hechos(a.pitchbook)
    comunes = sorted(set(fs) & set(pb))

    print("=" * 78)
    print("LOS DOS REPORTES DICEN LO MISMO?")
    print("  factsheet : %s" % os.path.basename(a.factsheet))
    print("  pitch book: %s" % os.path.basename(a.pitchbook))
    print("=" * 78)

    difieren = []
    for nombre in comunes:
        v1, o1 = fs[nombre]
        v2, o2 = pb[nombre]
        if valor_comparable(v1) != valor_comparable(v2):
            difieren.append((nombre, v1, o1, v2, o2))

    print("  %d datos compartidos | %d coinciden | %d difieren"
          % (len(comunes), len(comunes) - len(difieren), len(difieren)))

    if difieren:
        print("\n" + "-" * 78)
        print("NO COINCIDEN")
        print("-" * 78)
        for nombre, v1, o1, v2, o2 in difieren:
            print("  %-22s factsheet: %-30s (%s)" % (nombre, v1[:30], o1))
            print("  %-22s pitchbook: %-30s (%s)" % ("", v2[:30], o2))

    solo_fs = sorted(set(fs) - set(pb))
    solo_pb = sorted(set(pb) - set(fs))
    if solo_fs or solo_pb:
        print("\n  solo en el factsheet: %s" % (", ".join(solo_fs) or "—"))
        print("  solo en el pitch book: %s" % (", ".join(solo_pb) or "—"))

    print()
    if difieren:
        print("RESULTADO: %d dato(s) que los dos reportes cuentan distinto" % len(difieren))
        sys.exit(1)
    print("RESULTADO: los datos compartidos coinciden")
    sys.exit(0)


if __name__ == "__main__":
    main()
