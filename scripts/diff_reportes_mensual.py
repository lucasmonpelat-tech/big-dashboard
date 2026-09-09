# -*- coding: utf-8 -*-
"""Diff mes contra mes de los reportes de cliente del BIG (factsheet y pitch book).

Motivacion (2026-09-08): la columna "Inicio" de la tabla de rentabilidades del
factsheet quedo congelada en 8,75 / 8,99 desde Julio y se mando a clientes asi.
No fue un error de calculo: fue una celda que nadie toco y que nada obligaba a
tocar. Lo mismo habia pasado con las etiquetas del grafico de benchmark y con la
serie de calidad crediticia del AGG, congelada desde Abril.

Este script da vuelta la pregunta. En vez de buscar que esta mal, lista **todo lo
que quedo igual que el mes pasado** y obliga a justificarlo.

Compara tres cosas, por shape_id (estable, porque el deck del mes se copia del
anterior):
  - texto de cada autoshape / textbox
  - texto de cada celda de tabla
  - categorias, nombres de serie y valores de cada grafico

Solo marca como sospechoso lo que **contiene algun digito**: un titulo que no
cambia es normal, un numero que no cambia no lo es.

Lo que ya se sabe que es legitimamente estatico (ISIN, fee, minimo de inversion)
se declara en scripts/diff_reportes_ignore.json y deja de aparecer. Ese archivo
es el registro de "esto lo miramos y esta bien que no cambie".

Uso:
    python scripts/diff_reportes_mensual.py <actual.pptx> <anterior.pptx>
    python scripts/diff_reportes_mensual.py <actual> <anterior> --todo

Exit code 1 si quedo algo sin justificar, 0 si esta limpio.
"""
import argparse
import io
import json
import os
import re
import sys
import unicodedata

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

IGNORE_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                           "diff_reportes_ignore.json")
TIENE_DIGITO = re.compile(r"\d")
DEC_COMA = re.compile(r"\d,\d")
DEC_PUNTO = re.compile(r"\d\.\d")


def norm(s):
    """Normaliza para comparar: colapsa espacios en blanco."""
    s = unicodedata.normalize("NFKC", s or "")
    return re.sub(r"\s+", " ", s).strip()


def recolectar(path):
    """{clave: (etiqueta_legible, valor)} de todo el contenido textual y numerico."""
    pres = Presentation(path)
    out = {}

    def texto(slide_i, shape):
        if not shape.has_text_frame:
            return
        t = norm(shape.text_frame.text)
        if t:
            out[("txt", slide_i, shape.shape_id)] = (
                "s%d txt id=%d" % (slide_i, shape.shape_id), t)

    def tabla(slide_i, shape):
        for ri, row in enumerate(shape.table.rows):
            for ci, cell in enumerate(row.cells):
                t = norm(cell.text)
                if t:
                    out[("tbl", slide_i, shape.shape_id, ri, ci)] = (
                        "s%d tabla id=%d f%d c%d" % (slide_i, shape.shape_id, ri, ci), t)

    def grafico(slide_i, shape):
        ch = shape.chart
        sid = shape.shape_id
        try:
            cats = [norm(str(c)) for c in ch.plots[0].categories]
        except Exception:
            cats = []
        out[("cht_cat", slide_i, sid)] = (
            "s%d chart id=%d categorias" % (slide_i, sid), " | ".join(cats))
        for si, se in enumerate(ch.plots[0].series):
            nombre = norm(se.name or "")
            out[("cht_name", slide_i, sid, si)] = (
                "s%d chart id=%d serie %d nombre" % (slide_i, sid, si), nombre)
            for vi, v in enumerate(se.values):
                cat = cats[vi] if vi < len(cats) else ("#%d" % vi)
                out[("cht_val", slide_i, sid, si, vi)] = (
                    "s%d chart id=%d [%s] %s" % (slide_i, sid, nombre or si, cat),
                    "" if v is None else ("%.4f" % v))

    def walk(slide_i, shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(slide_i, sh.shapes)
                continue
            if sh.has_table:
                tabla(slide_i, sh)
            elif sh.has_chart:
                grafico(slide_i, sh)
            else:
                texto(slide_i, sh)

    for i, s in enumerate(pres.slides, start=1):
        walk(i, s.shapes)
    return out


def cargar_ignore(deck):
    if not os.path.exists(IGNORE_FILE):
        return {}
    return json.load(io.open(IGNORE_FILE, encoding="utf-8")).get(deck, {})


def separadores_mezclados(actual):
    """Slides donde unos pocos valores usan coma decimal y el resto punto.

    Esa es exactamente la huella que dejo el 'Inicio' arrastrado de Julio.
    """
    por_slide = {}
    for k, (etq, val) in actual.items():
        if not TIENE_DIGITO.search(val):
            continue
        slide = k[1]
        coma, punto = bool(DEC_COMA.search(val)), bool(DEC_PUNTO.search(val))
        if not (coma or punto):
            continue
        g = por_slide.setdefault(slide, {"coma": [], "punto": []})
        g["coma" if coma else "punto"].append((etq, val))
    res = []
    for slide, g in sorted(por_slide.items()):
        # minoria con coma rodeada de mayoria con punto (o al reves)
        if g["coma"] and g["punto"]:
            if len(g["coma"]) <= 3:
                res.append((slide, g["coma"], "coma", len(g["punto"])))
            elif len(g["punto"]) <= 3:
                res.append((slide, g["punto"], "punto", len(g["coma"])))
    return res


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("actual", help="pptx del mes que se esta cerrando")
    ap.add_argument("anterior", help="pptx del mes anterior")
    ap.add_argument("--deck", default=None,
                    help="clave del ignore list (default: se infiere del nombre)")
    ap.add_argument("--todo", action="store_true", help="listar tambien lo que cambio")
    a = ap.parse_args()

    deck = a.deck or ("factsheet" if "factsheet" in os.path.basename(a.actual).lower()
                      else "pitchbook")
    ign = cargar_ignore(deck)

    act, ant = recolectar(a.actual), recolectar(a.anterior)
    comunes = set(act) & set(ant)

    iguales, cambios = [], []
    for k in comunes:
        etq, va = act[k]
        vb = ant[k][1]
        (iguales if va == vb else cambios).append((etq, va, vb, k))
    iguales.sort(key=lambda x: (x[3][1], str(x[3])))
    cambios.sort(key=lambda x: (x[3][1], str(x[3])))

    con_numero = [x for x in iguales if TIENE_DIGITO.search(x[1])]
    sospechosos = [x for x in con_numero if x[0] not in ign]
    justificados = [x for x in con_numero if x[0] in ign]

    print("=" * 78)
    print("DIFF MENSUAL  [%s]" % deck)
    print("  actual   : %s" % os.path.basename(a.actual))
    print("  anterior : %s" % os.path.basename(a.anterior))
    print("=" * 78)
    print("  %d elementos comparables  |  %d cambiaron  |  %d quedaron igual"
          % (len(comunes), len(cambios), len(iguales)))
    print("  de los que quedaron igual, %d tienen numeros (%d ya justificados)"
          % (len(con_numero), len(justificados)))

    print("\n" + "-" * 78)
    print("A REVISAR — quedo IGUAL que el mes pasado y contiene numeros")
    print("-" * 78)
    if not sospechosos:
        print("  (nada: todos los numeros se movieron o estan justificados)")
    for etq, va, _vb, _k in sospechosos:
        print("  %-40s  %s" % (etq, va[:68]))

    mez = separadores_mezclados(act)
    if mez:
        print("\n" + "-" * 78)
        print("FORMATO — separador decimal distinto al de sus vecinos en la misma slide")
        print("-" * 78)
        for slide, raros, cual, n_otros in mez:
            for etq, val in raros:
                print("  %-40s %-20s usa %s, los otros %d de la slide no"
                      % (etq, val[:20], cual, n_otros))

    solo_act, solo_ant = set(act) - set(ant), set(ant) - set(act)
    if solo_act or solo_ant:
        print("\n" + "-" * 78)
        print("ESTRUCTURA — existe en un deck y no en el otro")
        print("-" * 78)
        for k in sorted(solo_act, key=str):
            print("  + solo actual    %-38s %s" % (act[k][0], act[k][1][:38]))
        for k in sorted(solo_ant, key=str):
            print("  - solo anterior  %-38s %s" % (ant[k][0], ant[k][1][:38]))

    if a.todo:
        print("\n" + "-" * 78)
        print("CAMBIOS")
        print("-" * 78)
        for etq, va, vb, _k in cambios:
            print("  %-40s  %-30s -> %s" % (etq, vb[:30], va[:30]))

    if justificados:
        print("\n  ya justificados en diff_reportes_ignore.json:")
        for etq, va, _vb, _k in justificados:
            print("     %-38s %-22s %s" % (etq, va[:22], ign[etq]))

    return 1 if sospechosos else 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
