# -*- coding: utf-8 -*-
"""Compara el NAV que publica Lynk contra el CAV oficial de ProCapital.

POR QUE EXISTE (2026-09-17)
---------------------------
El 14-Sep Lynk REESCRIBIO el NAV del 12-Ago-2026: lo cambio de 106.861, que
llevaba un mes publicado, a 103.215. Lo encontre de casualidad barriendo saltos
diarios mientras investigaba otra cosa. No habia nada en el repo capaz de
detectarlo.

La pregunta que quedo abierta fue la correcta: si reescribieron una fecha,
puede haber otras. Esto la contesta, y la sigue contestando todos los meses.

POR QUE CONTRA EL CAV Y NO CONTRA NUESTRO HISTORIAL DE GIT
----------------------------------------------------------
Comparar la serie de hoy contra la de hace un mes en git detecta "Lynk cambio
de opinion". Es util, pero no alcanza: un NAV que Lynk publico MAL desde el
primer dia nunca cambia, asi que esa comparacion no lo ve nunca.

El CAV es el registro del agente de calculo: la fuente de verdad del NAV de la
nota. Compararlo contra eso detecta las dos cosas -- lo reescrito y lo que
nacio mal.

MENSUAL, Y CORRE LOCAL
----------------------
Lucas sube el CAV una vez por mes para el proceso de fees, asi que la cadencia
mensual sale gratis: no hay dato nuevo entre medio que revisar.

Corre LOCAL, no en GitHub Actions, porque el CAV vive en Dropbox y el runner no
lo ve. Con --guardar se destila a data/cav_nav_oficial.json y queda versionado
en el repo: a partir de ahi el NAV oficial es parte de la capa de datos y no un
archivo suelto en una carpeta.

USO
---
    python scripts/check_lynk_vs_cav.py                 # compara y reporta
    python scripts/check_lynk_vs_cav.py --guardar       # ademas versiona el CAV
    python scripts/check_lynk_vs_cav.py --cav <ruta>    # un archivo puntual

Exit 1 si aparece una discrepancia NUEVA (las ya registradas en
data/lynk_puntos_malos.json no vuelven a alarmar: ya estan corregidas).
"""
from __future__ import annotations

import argparse
import datetime
import json
import os
import sys
import warnings
from pathlib import Path

warnings.filterwarnings("ignore", module="openpyxl")

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(Path(__file__).parent))

CARPETA_CAV = Path(r"C:\Users\lmonp\Dropbox\Maximus BIG\2026"
                   r"\Luiquidacion Comisiones BIG\Pro Capital Desgloze Fees")
SERIE_LYNK = ROOT / "data" / "lynk_nav_series.json"
CAV_DESTILADO = ROOT / "data" / "cav_nav_oficial.json"

HOJA = "NAV Calculation"
FILA_FECHA = 11
FILA_NAV = 57

# El CAV trae columnas PRE-CARGADAS para fechas futuras (llega hasta 2027) con
# un placeholder de -0.312. Sin este filtro, cada dia posterior al cierre del
# mes aparece como "discrepancia" gigante y el chequeo se vuelve ruido.
NAV_MIN, NAV_MAX = 10.0, 10000.0

# El CAV redondea a 3 decimales y Lynk publica con toda la precision del float.
# Medido sobre 296 fechas, la mayor diferencia por redondeo fue 0.0005.
TOLERANCIA = 0.002


def ultimo_cav(carpeta: Path = CARPETA_CAV) -> Path | None:
    """El CAV mas reciente de la carpeta de fees.

    Filtra por LS104 en el nombre a proposito: en esa carpeta conviven otros
    xlsx del proceso de fees (exports de posiciones, por ejemplo) y agarrar
    "el mas nuevo" a secas abria el archivo equivocado.
    """
    if not carpeta.exists():
        return None
    archivos = [p for p in carpeta.glob("*.xlsx")
                if not p.name.startswith("~$") and "LS104" in p.name.upper()]
    return max(archivos, key=lambda p: p.stat().st_mtime) if archivos else None


def leer_cav(path: Path) -> dict:
    """{fecha_iso: nav} del CAV, ya sin las columnas futuras sin cargar."""
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True)
    if HOJA not in wb.sheetnames:
        sys.exit(f"ERROR: {path.name} no tiene la hoja '{HOJA}'. "
                 f"Tiene: {wb.sheetnames}. Cambio el formato del CAV?")
    ws = wb[HOJA]
    out = {}
    for col in range(2, ws.max_column + 1):
        f = ws.cell(row=FILA_FECHA, column=col).value
        n = ws.cell(row=FILA_NAV, column=col).value
        if isinstance(f, datetime.datetime) and isinstance(n, (int, float)):
            if NAV_MIN <= n <= NAV_MAX:
                out[f.date().isoformat()] = float(n)
    return out


def conocidas() -> dict:
    """Discrepancias ya registradas y corregidas: no tienen que volver a alarmar."""
    try:
        from lynk_series import cargar_puntos_malos
        return cargar_puntos_malos()
    except Exception:
        return {}


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--cav", default=None, help="Ruta a un CAV puntual.")
    ap.add_argument("--guardar", action="store_true",
                    help="Versionar el CAV destilado en data/cav_nav_oficial.json.")
    ap.add_argument("--alerta", default=None, help="Donde dejar el JSON de alerta.")
    a = ap.parse_args()

    cav_path = Path(a.cav) if a.cav else ultimo_cav()
    if not cav_path or not cav_path.exists():
        sys.exit(f"ERROR: no encontre ningun CAV en {CARPETA_CAV}. "
                 f"Pasalo con --cav si esta en otro lado.")

    cav = leer_cav(cav_path)
    if not cav:
        sys.exit(f"ERROR: {cav_path.name} no tiene NAV utilizables en la fila {FILA_NAV}.")

    lynk = {p["date"]: p["value"]
            for p in json.loads(SERIE_LYNK.read_text(encoding="utf-8"))["series"]}
    comunes = sorted(set(cav) & set(lynk))

    print("=" * 74)
    print("  LYNK vs CAV OFICIAL (ProCapital LS104)")
    print("=" * 74)
    print(f"  CAV:  {cav_path.name}")
    print(f"        {len(cav)} fechas con NAV, {min(cav)} -> {max(cav)}")
    print(f"  Lynk: {len(lynk)} fechas, {min(lynk)} -> {max(lynk)}")
    print(f"  Comparadas: {len(comunes)}  (tolerancia {TOLERANCIA})")
    print()

    ya = conocidas()
    nuevas, viejas = [], []
    for d in comunes:
        dif = lynk[d] - cav[d]
        if abs(dif) <= TOLERANCIA:
            continue
        item = {"fecha": d, "lynk": round(lynk[d], 4),
                "cav_oficial": round(cav[d], 4), "diferencia": round(dif, 4)}
        (viejas if d in ya else nuevas).append(item)

    for i in viejas:
        print("  [conocida] %s  Lynk %.3f vs CAV %.3f  — ya corregida en el pipeline"
              % (i["fecha"], i["lynk"], i["cav_oficial"]))
    for i in nuevas:
        print("  [NUEVA]    %s  Lynk %.3f vs CAV %.3f  (%+.3f)"
              % (i["fecha"], i["lynk"], i["cav_oficial"], i["diferencia"]))

    if a.guardar:
        CAV_DESTILADO.write_text(json.dumps({
            "_doc": ("NAV oficial del agente de calculo (ProCapital, serie LS104). "
                     "Es la fuente de verdad del NAV de la nota. Se destila del CAV "
                     "que Lucas sube cada mes para el proceso de fees, para que el "
                     "registro oficial viva en el repo y no solo en Dropbox."),
            # refreshedAt y no "generado": es la clave que ya lee
            # check_race_freshness.py, que es quien avisa si este archivo
            # quedo viejo porque no se corrio el chequeo del mes.
            "refreshedAt": datetime.date.today().isoformat(),
            "cubre_hasta": max(cav),
            "origen": cav_path.name,
            "isin": "XS3037627794",
            "series_number": "LS104",
            "series": [{"date": d, "value": cav[d]} for d in sorted(cav)],
        }, indent=2, ensure_ascii=False), encoding="utf-8")
        print(f"\n  CAV destilado -> {CAV_DESTILADO.relative_to(ROOT)} ({len(cav)} fechas)")

    print()
    if not nuevas:
        print("  Sin discrepancias nuevas.")
        if viejas:
            print("  (%d conocida(s), ya corregida(s) — sacar de lynk_puntos_malos.json "
                  "cuando Lynk las arregle)" % len(viejas))
        sys.exit(0)

    print("  %d discrepancia(s) NUEVA(S)." % len(nuevas))
    if a.alerta:
        os.makedirs(os.path.dirname(a.alerta), exist_ok=True)
        with open(a.alerta, "w", encoding="utf-8") as f:
            json.dump({
                "date": datetime.date.today().isoformat(),
                "tipo": "lynk_vs_cav_discrepancia",
                "issues": ["%s: Lynk dice %.3f y el CAV oficial %.3f (%+.3f)"
                           % (i["fecha"], i["lynk"], i["cav_oficial"], i["diferencia"])
                           for i in nuevas],
                "accion": ("El NAV que publica Lynk no coincide con el registro del "
                           "agente de calculo. El CAV manda. Agregar la fecha a "
                           "data/lynk_puntos_malos.json con valor_correcto = el del "
                           "CAV, y reclamarle a Lynk."),
                "cav_usado": cav_path.name,
                "detalle": nuevas,
            }, f, indent=2, ensure_ascii=False)
        print("  Alerta escrita en %s" % a.alerta)
    sys.exit(1)


if __name__ == "__main__":
    main()
